from __future__ import annotations
import json, re
from pathlib import Path
from .models import Settings, ScenePlan, LogFn

def split_sentences(text: str, max_words: int = 12) -> list[str]:
    text=re.sub(r"\s+"," ",text).strip()
    if not text:return []
    sentences=re.split(r"(?<=[.!?…])\s+",text)
    out=[]
    for sentence in sentences:
        words=sentence.split()
        while len(words)>max_words:
            cut=max_words
            # prefer a comma/semicolon boundary near the end
            for j in range(max_words-1,max(3,max_words//2)-1,-1):
                if re.search(r"[,;:，；：]$",words[j-1]): cut=j; break
            out.append(" ".join(words[:cut]).strip()); words=words[cut:]
        if words: out.append(" ".join(words).strip())
    return [x for x in out if x]

def balance_text_lines(text: str, max_chars: int = 35) -> str:
    words=text.split()
    if len(text)<=max_chars or len(words)<2:return text
    best=None
    for i in range(1,len(words)):
        a=" ".join(words[:i]); b=" ".join(words[i:])
        score=max(len(a),len(b))+abs(len(a)-len(b))*0.25
        if len(a)<=max_chars*1.35 and len(b)<=max_chars*1.35 and (best is None or score<best[0]): best=(score,a,b)
    return f"{best[1]}\n{best[2]}" if best else text

def _manifest_texts(path: Path) -> dict[int,str]:
    data=json.loads(path.read_text(encoding="utf-8"))
    out={}
    def add(n,v):
        try:n=int(n)
        except Exception:return
        if isinstance(v,str) and v.strip(): out[n]=v.strip(); return
        if isinstance(v,dict):
            for k in ("script","text","content","narration","voiceover"):
                if isinstance(v.get(k),str) and v[k].strip(): out[n]=v[k].strip(); return
    if isinstance(data,dict):
        if isinstance(data.get("scenes"),list):
            for item in data["scenes"]:
                if isinstance(item,dict): add(item.get("scene_id",item.get("scene",item.get("id",item.get("number")))),item)
        else:
            for k,v in data.items(): add(k,v)
    elif isinstance(data,list):
        for item in data:
            if isinstance(item,dict): add(item.get("scene_id",item.get("scene",item.get("id",item.get("number")))),item)
    return out

# --------------------------------------------------------------------------- #
# Excel lời thoại
#
# Mặc định (file phân cảnh chuẩn): cột A = số thứ tự cảnh, cột B = lời thoại,
# có hoặc không có dòng tiêu đề đều được. Nếu file có dòng tiêu đề rõ ràng
# (vd "STT | Lời thoại") thì tool tìm cột theo tên; người dùng cũng có thể gõ
# tên cột, chữ cái cột ("B") hoặc số thứ tự cột ("2") vào ô Cột lời thoại /
# Cột số cảnh để chỉ định thẳng.
# --------------------------------------------------------------------------- #
SCRIPT_TERMS=("script","text","content","narration","voice","thoại","thoai","kịch","kich",
              "lời","loi","nội dung","noi dung","dialog")
SCENE_TERMS=("scene","panel","stt","number","cảnh","canh","số","so","#","index","no.","id")
DEFAULT_SCENE_COL=0     # cột A
DEFAULT_SCRIPT_COL=1    # cột B


def _is_number(v) -> bool:
    if isinstance(v,(int,float)) and not isinstance(v,bool): return True
    try: float(str(v).strip()); return True
    except (TypeError,ValueError): return False


def _looks_like_header(row) -> bool:
    """Dòng tiêu đề thật: ô ngắn, ô đầu không phải số, có ít nhất một tên cột quen."""
    vals=[str(v).strip() for v in row if v not in (None,"") and str(v).strip()]
    if not vals or len(vals)>30: return False
    if row and row[0] not in (None,"") and _is_number(row[0]): return False
    if any(len(v)>40 for v in vals): return False
    low=[v.lower() for v in vals]
    terms=SCRIPT_TERMS+SCENE_TERMS
    return any(v==t or v.startswith(t) or t in v.split() for v in low for t in terms)


def _find_header(rows) -> int:
    """Chỉ số dòng tiêu đề, hoặc -1 nếu file không có tiêu đề (dữ liệu từ dòng 1)."""
    for i,row in enumerate(rows[:5]):
        if row and _looks_like_header(row): return i
        if row and any(v not in (None,"") for v in row):
            # dòng dữ liệu đầu tiên đã xuất hiện mà chưa thấy tiêu đề → không có
            if _is_number(row[0]) or (len(row)>1 and len(str(row[1] or ""))>40): return -1
    return -1


def _col_letter(i: int) -> str:
    s=""; i+=1
    while i: i,r=divmod(i-1,26); s=chr(65+r)+s
    return s


def _resolve_col(spec: str, headers: list[str]) -> int | None:
    """'B' / 'b' → 1, '2' → 1, tên tiêu đề → vị trí của nó, '' → None."""
    spec=(spec or "").strip()
    if not spec: return None
    for i,h in enumerate(headers):
        if h and h.strip().lower()==spec.lower(): return i
    if spec.isdigit() and int(spec)>=1: return int(spec)-1
    m=re.fullmatch(r"(?i)(?:cột\s*|col(?:umn)?\s*)?([a-z]{1,2})",spec)
    if m:
        n=0
        for ch in m.group(1).upper(): n=n*26+ord(ch)-64
        return n-1
    m=re.match(r"(?i)^cột\s+([a-z]{1,2})\b",spec)
    if m: return _resolve_col(m.group(1),[])
    return None


def _guess_col(headers: list[str], terms, exclude=None) -> int | None:
    for i,h in enumerate(headers):
        if i==exclude: continue
        hl=h.lower()
        if hl and any(hl==t or hl.startswith(t) or t in hl.split() for t in terms): return i
    for i,h in enumerate(headers):
        if i==exclude: continue
        hl=h.lower()
        if hl and any(t in hl for t in terms if len(t)>2): return i
    return None


def excel_columns(path: Path, sheet_name: str="") -> tuple[list[str],list[str]]:
    """(danh sách sheet, tên các cột). File không có tiêu đề → 'Cột A', 'Cột B'…"""
    import openpyxl
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        sheets=wb.sheetnames
        ws=wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
        rows=list(ws.iter_rows(min_row=1,max_row=20,values_only=True))
    finally:
        wb.close()
    if not rows: return sheets,[]
    idx=_find_header(rows)
    if idx>=0:
        return sheets,[str(v or "").strip() or f"Cột {_col_letter(i)}" for i,v in enumerate(rows[idx])]
    width=max(len(r) for r in rows)
    return sheets,[f"Cột {_col_letter(i)}" for i in range(width)]


def guess_excel_columns(headers: list[str]) -> tuple[str,str]:
    """(cột lời thoại, cột số cảnh) để điền sẵn vào giao diện."""
    if headers and all(h.startswith("Cột ") for h in headers):
        return (headers[DEFAULT_SCRIPT_COL] if len(headers)>1 else "",
                headers[DEFAULT_SCENE_COL] if headers else "")
    sidx=_guess_col(headers,SCRIPT_TERMS)
    nidx=_guess_col(headers,SCENE_TERMS,exclude=sidx)
    if sidx is None and len(headers)>1: sidx=DEFAULT_SCRIPT_COL
    if nidx is None and headers and sidx!=DEFAULT_SCENE_COL: nidx=DEFAULT_SCENE_COL
    return (headers[sidx] if sidx is not None else "",
            headers[nidx] if nidx is not None else "")


def _excel_texts(path: Path, script_col: str, scene_col: str="", sheet_name: str="") -> dict[int,str]:
    import openpyxl
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
        rows=list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    rows=[r for r in rows if r is not None]
    if not rows: raise ValueError("File Excel trống")
    hidx=_find_header(rows)
    headers=[str(v or "").strip() for v in rows[hidx]] if hidx>=0 else []
    width=max(len(r) for r in rows)
    if not headers: headers=[f"Cột {_col_letter(i)}" for i in range(width)]

    sidx=_resolve_col(script_col,headers)
    nidx=_resolve_col(scene_col,headers)
    if sidx is None:
        sidx=_guess_col(headers,SCRIPT_TERMS) if hidx>=0 else None
        if sidx is None: sidx=DEFAULT_SCRIPT_COL
    if nidx is None:
        nidx=_guess_col(headers,SCENE_TERMS,exclude=sidx) if hidx>=0 else None
        if nidx is None and sidx!=DEFAULT_SCENE_COL: nidx=DEFAULT_SCENE_COL
    if sidx>=width and script_col:
        sidx=DEFAULT_SCRIPT_COL           # tên cột cũ không còn khớp → dùng mặc định
    if nidx is not None and nidx>=width:
        nidx=DEFAULT_SCENE_COL if sidx!=DEFAULT_SCENE_COL else None
    if sidx>=width:
        raise ValueError(f"Excel chỉ có {width} cột — không có cột lời thoại "
                         f"{_col_letter(sidx)}. Mặc định: cột A = số cảnh, cột B = lời thoại.")

    out={}; auto=1
    for row in rows[hidx+1:]:
        if sidx>=len(row):continue
        txt=str(row[sidx] if row[sidx] is not None else "").strip()
        if not txt:continue
        n=auto
        if nidx is not None and nidx<len(row) and row[nidx] not in (None,""):
            try:n=int(float(str(row[nidx]).strip()))
            except Exception:continue    # dòng ghi chú / tiêu đề phụ → bỏ qua
        out[n]=txt; auto+=1
    if not out:
        raise ValueError("Không đọc được dòng lời thoại nào từ Excel. Mặc định: cột A = số "
                         "cảnh, cột B = lời thoại; hoặc gõ tên/chữ cái cột ở Cài đặt → Phụ đề.")
    return out

def load_scene_texts(s: Settings, plan, log: LogFn=print) -> dict[int,str]:
    """`plan` nhận list[ScenePlan] hoặc thẳng list số cảnh."""
    src=s.subtitle_source.lower()
    nums={p if isinstance(p,int) else p.number for p in plan}; out={}
    if src=="off": return {}
    if src in ("manifest","auto"):
        candidates=[s.input_dir/"_manifest.json", s.input_dir.parent/"_manifest.json"]
        mp=next((p for p in candidates if p.is_file()),None)
        if mp:
            out=_manifest_texts(mp); log(f"[SUB] Manifest: {mp.name} ({len(out)} scene)")
            if src=="manifest": return {k:v for k,v in out.items() if k in nums}
    if src in ("excel","auto") and s.subtitle_file and Path(s.subtitle_file).is_file():
        p=Path(s.subtitle_file)
        if p.suffix.lower() in (".xlsx",".xlsm"):
            out=_excel_texts(p,s.excel_script_col,s.excel_scene_col,s.excel_sheet); log(f"[SUB] Excel: {p.name} ({len(out)} scene)")
            if src=="excel" or out:return {k:v for k,v in out.items() if k in nums}
    if src in ("texts","auto"):
        td=s.input_dir/"Texts"
        if td.is_dir():
            for f in td.glob("*.txt"):
                if f.stem.isdigit(): out[int(f.stem)]=f.read_text(encoding="utf-8-sig").strip()
            if out: log(f"[SUB] Texts/: {len(out)} file")
    return {k:v for k,v in out.items() if k in nums}

def srt_ts(sec: float) -> str:
    ms=max(0,int(round(sec*1000))); h,ms=divmod(ms,3600000); m,ms=divmod(ms,60000); ss,ms=divmod(ms,1000)
    return f"{h:02d}:{m:02d}:{ss:02d},{ms:03d}"

def build_srt(s: Settings, plan: list[ScenePlan], out_path: Path, log: LogFn=print) -> int:
    texts=load_scene_texts(s,plan,log); entries=[]
    for p in plan:
        raw=texts.get(p.number,"").strip()
        if not raw:
            log(f"[WARN] {p.number:04d} audio không có lời thoại")
            continue
        chunks=split_sentences(raw,s.sub_max_words) or [raw]
        weights=[max(1,len(c.split())) for c in chunks]; total=sum(weights); t=p.start
        for idx,(c,w) in enumerate(zip(chunks,weights)):
            dur=p.audio_duration*w/total
            # keep readability, but preserve exact scene boundary by adjusting last chunk
            if idx==len(chunks)-1: dur=p.audio_end-t
            entries.append((t,t+dur,balance_text_lines(c))); t+=dur
    out_path.parent.mkdir(parents=True,exist_ok=True)
    with out_path.open("w",encoding="utf-8") as f:
        for i,(st,en,txt) in enumerate(entries,1):
            f.write(f"{i}\n{srt_ts(st)} --> {srt_ts(en)}\n{txt}\n\n")
    log(f"[BUILD] Subtitles: {len(entries)} câu → {out_path}")
    return len(entries)
